#!/usr/bin/env python3
"""Unit tests for structured Excel extraction and schema mapping."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


class DummyOpenAIClient:
    def __init__(self, **kwargs):
        self.kwargs = kwargs


class StructuredExtractionTests(unittest.TestCase):
    def setUp(self):
        self.work_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.work_dir.cleanup)
        self.excel_path = Path(self.work_dir.name) / "book.xlsx"

        wb = Workbook()

        ws1 = wb.active
        ws1.title = "Tabular"
        ws1.append(["control_id", "audit_procedure", "test_steps", "test_result", "conclusion"])
        ws1.append(["CTRL-001", "抽样检查审批", "核对审批流并复核日志", "发现1个例外", "控制基本有效"])
        ws1.merge_cells("A4:B4")
        ws1["A4"] = "Merged Note"

        ws2 = wb.create_sheet("KeyValue")
        ws2.append(["控制编号", "CTRL-100"])
        ws2.append(["审计程序", "检查审批、复核日志、追踪异常处理"])
        ws2.append(["测试步骤", "抽样2个月交易并追踪例外整改"])
        ws2.append(["测试结果", "发现审批缺失2笔"])
        ws2.append(["结论", "控制无效"])

        ws_gap = wb.create_sheet("KeyValueGap")
        ws_gap["A1"] = "控制编号"
        ws_gap["C1"] = "CTRL-300"
        ws_gap["A2"] = "执行有效性测试结论"
        ws_gap["C2"] = "无效"

        ws_split = wb.create_sheet("KeyValueSplit")
        ws_split["A1"] = "控制编号"
        ws_split["C1"] = "CTRL-400"
        ws_split["A3"] = "执行有效性测试结论"
        ws_split["C3"] = "有效"
        ws_split["A5"] = "是否发现异常"
        ws_split["C5"] = "No"

        ws3 = wb.create_sheet("OffsetHeader")
        ws3.append(["底稿封面"])
        ws3.append(["步骤说明"])
        ws3.append(["序号", "控制编号", "控制描述", "执行有效性测试结论", "是否发现异常"])
        ws3.append([1, "CTRL-200", "账号权限审批", "有效", "N"])
        ws3.append([2, "", "离职账号禁用及时", "无效", "Y"])
        ws3.append([3, "CTRL-201", "日志审计", "有效", "N"])

        ws_qa = wb.create_sheet("QnAOnly")
        ws_qa.append(["问题", "回答"])
        ws_qa.append(["对接口控制执行了哪些测试？", "检查接口日志并核对异常处理，未见异常。"])
        ws_qa.append(["项目组还执行了哪些程序？", "抽样核对3笔数据传输记录，结果一致。"])

        ws_step = wb.create_sheet("StepBlock")
        ws_step["A1"] = "控制编号"
        ws_step["B1"] = "CTRL-500"
        ws_step["A2"] = "控制活动"
        ws_step["B2"] = "特权访问控制"
        ws_step["A4"] = "测试步骤"
        ws_step["B4"] = "执行有效性标准审计程序"
        ws_step["C4"] = "执行有效性执行的审计程序"
        ws_step["A5"] = "1"
        ws_step["B5"] = "检查审批记录"
        ws_step["C5"] = "访谈管理员并检查审批单据"
        ws_step["A6"] = "2"
        ws_step["B6"] = "检查职责分离"
        ws_step["C6"] = "通过系统角色清单核对，发现1例冲突"
        ws_step["A7"] = "执行有效性测试结论"
        ws_step["C7"] = "无效"
        ws_step["A9"] = "是否发现异常"
        ws_step["B9"] = "Yes"

        wb.save(self.excel_path)

    def _build_reviewer(self):
        from excel_image_review import ExcelImageReviewer

        return ExcelImageReviewer(str(self.excel_path), output_dir=self.work_dir.name)

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_extract_sheet_structure_preserves_coordinates_and_merged_ranges(self):
        reviewer = self._build_reviewer()
        structure = reviewer.extract_sheet_structure("Tabular")

        self.assertEqual(structure["sheet_name"], "Tabular")
        self.assertGreater(structure["non_empty_cell_count"], 0)
        self.assertTrue(any(cell["coord"] == "A2" for cell in structure["cells"]))
        self.assertTrue(any(item["range"] == "A4:B4" for item in structure["merged_ranges"]))
        self.assertTrue(structure["table_regions"])

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_from_tabular_headers(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("Tabular")
        records = reviewer.map_sheet_to_schema("Tabular")

        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record["control_id"], "CTRL-001")
        self.assertEqual(record["audit_procedure"], "抽样检查审批")
        self.assertEqual(record["test_steps"], "核对审批流并复核日志")
        self.assertEqual(record["test_result"], "发现1个例外")
        self.assertEqual(record["conclusion"], "控制基本有效")
        self.assertEqual(record["source_cells"]["control_id"], "A2")

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_from_key_value_layout(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("KeyValue")
        records = reviewer.map_sheet_to_schema("KeyValue")

        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record["control_id"], "CTRL-100")
        self.assertIn("复核日志", record["audit_procedure"])
        self.assertIn("审批缺失", record["test_result"])
        self.assertEqual(record["conclusion"], "控制无效")

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_from_key_value_with_gap_column(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("KeyValueGap")
        records = reviewer.map_sheet_to_schema("KeyValueGap")

        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record["control_id"], "CTRL-300")
        self.assertEqual(record["conclusion"], "无效")

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_keeps_split_key_value_regions_with_complete_fields(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("KeyValueSplit")
        records = reviewer.map_sheet_to_schema("KeyValueSplit")

        # The fixture intentionally separates each key/value row with blank rows,
        # so region detection creates 3 independent KV regions.
        self.assertEqual(len(records), 3)
        self.assertTrue(all(record["control_id"] == "CTRL-400" for record in records))
        self.assertTrue(any(record["conclusion"] == "有效" for record in records))
        self.assertTrue(any(record["exception_flag"] == "No" for record in records))

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_detects_offset_header_and_propagates_context(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("OffsetHeader")
        records = reviewer.map_sheet_to_schema("OffsetHeader")

        self.assertEqual(len(records), 3)

        self.assertEqual(records[0]["control_id"], "CTRL-200")
        self.assertEqual(records[0]["control_description"], "账号权限审批")
        self.assertEqual(records[0]["conclusion"], "有效")
        self.assertEqual(records[0]["exception_flag"], "N")

        self.assertEqual(records[1]["control_id"], "CTRL-200")
        self.assertEqual(records[1]["control_description"], "离职账号禁用及时")
        self.assertEqual(records[1]["conclusion"], "无效")
        self.assertEqual(records[1]["exception_flag"], "Y")

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_fallbacks_to_question_answer_layout(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("QnAOnly")
        records = reviewer.map_sheet_to_schema("QnAOnly")

        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertIn("接口控制", record["audit_procedure"])
        self.assertIn("检查接口日志", record["test_steps"])

    @patch("excel_image_review.OpenAI", new=DummyOpenAIClient)
    def test_map_sheet_to_schema_extracts_step_block_details(self):
        reviewer = self._build_reviewer()
        reviewer.extract_sheet_structure("StepBlock")
        records = reviewer.map_sheet_to_schema("StepBlock")

        self.assertTrue(any("访谈管理员并检查审批单据" in (record.get("test_steps") or "") for record in records))
        self.assertTrue(any("检查审批记录" in (record.get("audit_procedure") or "") for record in records))
        self.assertTrue(any(record.get("conclusion") == "无效" for record in records))
        self.assertTrue(any(record.get("exception_flag") == "Yes" for record in records))


if __name__ == "__main__":
    unittest.main()
